from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def create_architecture_doc():
    document = Document()
    
    # Style configuration
    style = document.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(13)

    # Title
    title = document.add_heading('GIẢI THÍCH KIẾN TRÚC HỆ THỐNG', level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    document.add_paragraph()

    # Introduction
    document.add_paragraph(
        'Hệ thống PharmaCare được xây dựng dựa trên kiến trúc MERN Stack (MongoDB, Express, React, Node.js) '
        'kết hợp với mô hình Client-Server. Kiến trúc được chia thành 3 tầng chính (3-Tier Architecture): '
        'Client Tier, Application Server Tier, và Data & External Tier.'
    )

    # Insert Diagram
    document.add_heading('1. Sơ đồ kiến trúc tổng quan', level=2)
    try:
        document.add_picture('architecture_diagram_generated.png', width=Cm(16))
        last_p = document.paragraphs[-1]
        last_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption = document.add_paragraph('Hình 1: Kiến trúc chi tiết hệ thống PharmaCare', style='Caption')
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    except Exception as e:
        document.add_paragraph(f'[Không tìm thấy hình ảnh sơ đồ: {e}]')

    document.add_paragraph()

    # Detailed Explanation
    document.add_heading('2. Chi tiết các thành phần', level=2)

    # 2.1 Client Tier
    document.add_heading('2.1. Client Tier (Tầng Giao Diện)', level=3)
    document.add_paragraph('Chịu trách nhiệm tương tác trực tiếp với người dùng. Gồm 2 ứng dụng riêng biệt:')
    
    p = document.add_paragraph()
    p.add_run('• Customer App: ').bold = True
    p.add_run('Xây dựng bằng React.js + Vite. Đây là website dành cho khách hàng tìm kiếm thuốc, đặt hàng online. ')
    p.add_run('Sử dụng Redux/Context API để quản lý trạng thái (Giỏ hàng, Auth) và Axios để gọi API.')

    p = document.add_paragraph()
    p.add_run('• Admin Panel: ').bold = True
    p.add_run('Xây dựng bằng React.js + TypeScript + Shadcn/UI. Dành cho nhân viên/quản trị viên quản lý kho, đơn hàng, báo cáo. ')
    p.add_run('Tập trung vào tính bảo mật, xử lý dữ liệu phức tạp và hiển thị biểu đồ.')

    # 2.2 Server Tier
    document.add_heading('2.2. Application Server Tier (Tầng Ứng Dụng)', level=3)
    document.add_paragraph('Trung tâm xử lý logic của hệ thống, xây dựng trên nền tảng Node.js và Express.js. Được chia thành các lóp (Layered Architecture):')

    p = document.add_paragraph()
    p.add_run('• Middleware Layer: ').bold = True
    p.add_run('Lớp bảo vệ đầu tiên. Xử lý xác thực người dùng (Auth/JWT), kiểm tra dữ liệu đầu vào (Validation/Joi), bảo mật (Helmet/CORS) và xử lý lỗi tập trung.')

    p = document.add_paragraph()
    p.add_run('• Controller Layer: ').bold = True
    p.add_run('Tiếp nhận request từ Client, điều phối xử lý và trả về response. Chia nhỏ theo các module: Auth, Product, Order, User, Report.')

    p = document.add_paragraph()
    p.add_run('• Service Layer: ').bold = True
    p.add_run('Chứa Business Logic cốt lõi (Vd: Tính toán tồn kho FEFO, Xử lý thanh toán VNPay). Tách biệt logic kinh doanh khỏi Controller giúp code dễ bảo trì.')

    # 2.3 Data Tier
    document.add_heading('2.3. Data & External Tier (Tầng Dữ Liệu & Bên Ngoài)', level=3)
    
    p = document.add_paragraph()
    p.add_run('• Database Server (MongoDB): ').bold = True
    p.add_run('Cơ sở dữ liệu NoSQL hướng tài liệu (Document-oriented). Sử dụng Mongoose ODM để mô hình hóa dữ liệu (User, Product, Batch, Invoice). Phù hợp với dữ liệu bán hàng linh động.')

    p = document.add_paragraph()
    p.add_run('• External Services: ').bold = True
    p.add_run('Các dịch vụ bên thứ 3 tích hợp:')
    document.add_paragraph('  - VNPay Gateway: Cổng thanh toán trực tuyến.')
    document.add_paragraph('  - SMTP Server: Gửi email xác thực, thông báo đơn hàng.')

    # 3. Flow
    document.add_heading('3. Luồng hoạt động chính', level=2)
    document.add_paragraph('1. Client gửi HTTP Request (JSON) tới REST API.')
    document.add_paragraph('2. Server đi qua Middleware (Auth/Validation).')
    document.add_paragraph('3. Controller gọi Service để xử lý nghiệp vụ.')
    document.add_paragraph('4. Service tương tác với Database (Mongoose) hoặc gọi External API (VNPay).')
    document.add_paragraph('5. Server trả về Response cho Client hiển thị.')

    # Save
    file_path = 'Giai_Thich_Kien_Truc.docx'
    document.save(file_path)
    print(f"Successfully created {file_path}")

if __name__ == '__main__':
    create_architecture_doc()
