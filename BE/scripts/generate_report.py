from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE

def create_report():
    document = Document()
    
    # Configure default style
    style = document.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(13)
    
    # ========== TITLE PAGE ==========
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run('BÁO CÁO ĐỒ ÁN\n')
    run.bold = True
    run.font.size = Pt(18)
    
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run2 = subtitle.add_run('WEBSITE BÁN THUỐC TRỰC TUYẾN PHARMACARE')
    run2.bold = True
    run2.font.size = Pt(16)
    
    document.add_page_break()
    
    # ========== CHAPTER 1 ==========
    chapter1 = document.add_heading('CHƯƠNG 1: GIỚI THIỆU ĐỀ TÀI', level=1)
    
    # 1.1 Lý do chọn đề tài
    document.add_heading('1.1. Lý do chọn đề tài', level=2)
    document.add_paragraph(
        'Trong bối cảnh xã hội hiện đại, nhu cầu mua sắm trực tuyến ngày càng tăng cao, '
        'đặc biệt là trong lĩnh vực y tế và dược phẩm. Đại dịch COVID-19 đã thúc đẩy xu hướng '
        'mua thuốc online, giúp người dân tiếp cận thuốc men một cách nhanh chóng và an toàn.'
    )
    document.add_paragraph(
        'Tuy nhiên, các hệ thống bán thuốc trực tuyến hiện tại tại Việt Nam vẫn còn nhiều hạn chế:'
    )
    document.add_paragraph('• Thiếu tính năng quản lý kho hàng đa chi nhánh', style='List Bullet')
    document.add_paragraph('• Chưa hỗ trợ quản lý hạn sử dụng (expiry date) của thuốc', style='List Bullet')
    document.add_paragraph('• Giao diện người dùng chưa thân thiện', style='List Bullet')
    document.add_paragraph('• Thiếu tích hợp thanh toán trực tuyến (VNPay, Momo...)', style='List Bullet')
    document.add_paragraph(
        'Xuất phát từ những lý do trên, nhóm chúng em quyết định xây dựng hệ thống '
        '"PharmaCare - Website Bán Thuốc Trực Tuyến" nhằm giải quyết các vấn đề nêu trên.'
    )
    
    # 1.2 Mục tiêu
    document.add_heading('1.2. Mục tiêu', level=2)
    document.add_paragraph('Mục tiêu tổng quát:', style='Normal').runs[0].bold = True
    document.add_paragraph(
        'Xây dựng một hệ thống website bán thuốc trực tuyến hoàn chỉnh, bao gồm cả Frontend '
        'cho khách hàng và Admin Panel cho quản trị viên/nhân viên.'
    )
    document.add_paragraph('Mục tiêu cụ thể:', style='Normal').runs[0].bold = True
    document.add_paragraph('1. Xây dựng giao diện Frontend hiện đại, responsive cho khách hàng mua hàng online', style='List Number')
    document.add_paragraph('2. Xây dựng Admin Panel quản lý: Sản phẩm, Đơn hàng, Kho hàng, Nhân viên, Khách hàng', style='List Number')
    document.add_paragraph('3. Tích hợp thanh toán VNPay', style='List Number')
    document.add_paragraph('4. Quản lý kho hàng đa chi nhánh (Central + Branch)', style='List Number')
    document.add_paragraph('5. Quản lý lô hàng (batch) với hạn sử dụng', style='List Number')
    document.add_paragraph('6. Hỗ trợ bán hàng tại quầy (POS) cho nhân viên chi nhánh', style='List Number')
    
    # 1.3 Phạm vi thực hiện
    document.add_heading('1.3. Phạm vi thực hiện', level=2)
    document.add_paragraph('Phạm vi chức năng:', style='Normal').runs[0].bold = True
    
    # Table for scope
    table = document.add_table(rows=1, cols=2)
    table.style = 'Table Grid'
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'Thành phần'
    hdr_cells[1].text = 'Chức năng'
    hdr_cells[0].paragraphs[0].runs[0].font.bold = True
    hdr_cells[1].paragraphs[0].runs[0].font.bold = True
    
    data = [
        ['Frontend Website', 'Đăng ký/Đăng nhập, Xem sản phẩm, Giỏ hàng, Thanh toán (COD/VNPay), Theo dõi đơn hàng'],
        ['Admin Panel', 'Quản lý sản phẩm, Quản lý đơn hàng, Quản lý kho, Nhập hàng, Chuyển kho, Quản lý nhân viên/khách hàng'],
        ['Backend API', 'RESTful API với Node.js + Express, MongoDB database'],
        ['Tích hợp', 'VNPay Payment Gateway']
    ]
    for item in data:
        row_cells = table.add_row().cells
        row_cells[0].text = item[0]
        row_cells[1].text = item[1]
    
    document.add_paragraph()
    document.add_paragraph('Phạm vi công nghệ:', style='Normal').runs[0].bold = True
    document.add_paragraph('• Frontend: React.js + TypeScript + Vite', style='List Bullet')
    document.add_paragraph('• Admin: React.js + TypeScript + Shadcn/UI', style='List Bullet')
    document.add_paragraph('• Backend: Node.js + Express.js', style='List Bullet')
    document.add_paragraph('• Database: MongoDB (NoSQL)', style='List Bullet')
    document.add_paragraph('• Payment: VNPay Sandbox', style='List Bullet')

    document.add_heading('1.3.1. Mô hình Phân phối & Logistic (Logistics Model)', level=3)
    document.add_paragraph(
        '  Hệ thống vận hành theo mô hình Chuỗi cung ứng tập trung:', style='Normal'
    )
    document.add_paragraph('• Kho Tổng (Central): Nhập hàng từ Nhà cung cấp, lưu trữ số lượng lớn.', style='List Bullet')
    document.add_paragraph('• Điều chuyển (Transfer): Phân bổ hàng hóa từ Kho Tổng về các Kho Chi nhánh.', style='List Bullet')
    document.add_paragraph('• Kho Chi nhánh (Branch): Phục vụ bán hàng tại chỗ và giao hàng Online chặng cuối (Last-mile delivery).', style='List Bullet')
    
    # 1.4 Bố cục báo cáo
    document.add_heading('1.4. Bố cục báo cáo', level=2)
    document.add_paragraph('Báo cáo được chia thành các chương sau:')
    document.add_paragraph('• Chương 1: Giới thiệu đề tài - Trình bày lý do, mục tiêu và phạm vi', style='List Bullet')
    document.add_paragraph('• Chương 2: Cơ sở lý thuyết - Các công nghệ sử dụng (React, Node.js, MongoDB...)', style='List Bullet')
    document.add_paragraph('• Chương 3: Phân tích thiết kế - Sơ đồ Use Case, ERD, Database Schema', style='List Bullet')
    document.add_paragraph('• Chương 4: Hiện thực - Demo các chức năng chính', style='List Bullet')
    document.add_paragraph('• Chương 5: Kết luận - Đánh giá kết quả, hướng phát triển', style='List Bullet')

    document.add_page_break()

    # ========== CHAPTER 3 ==========
    document.add_heading('CHƯƠNG 3: PHÂN TÍCH VÀ THIẾT KẾ HỆ THỐNG', level=1)

    # 3.1 Architecture
    document.add_heading('3.1. Kiến trúc hệ thống', level=2)
    document.add_paragraph(
        'Hệ thống được xây dựng theo kiến trúc Client-Server với mô hình MERN Stack (MongoDB - Express - React - Node), '
        'chia tách rõ ràng giữa Frontend và Backend:'
    )
    document.add_paragraph('• Frontend (Client Side):', style='List Bullet')
    document.add_paragraph(
        '  Xây dựng bằng React.js, bao gồm hai phân hệ riêng biệt: Website bán hàng (cho khách hàng) và '
        'Admin Panel (cho nhân viên). Giao tiếp với Backend thông qua RESTful API.', style='Normal'
    )

    # Insert Architecture Diagram
    # Insert Architecture Diagram
    try:
        document.add_picture('architecture_diagram_generated.png', width=Cm(16)) # Increased width slightly
        last_paragraph = document.paragraphs[-1] 
        last_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption = document.add_paragraph('Hình 3.1: Sơ đồ kiến trúc hệ thống MERN Stack', style='Caption')
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    except Exception as e:
        print(f"Warning: Could not add image. Error: {e}")
        p = document.add_paragraph('[SƠ ĐỒ KIẾN TRÚC HỆ THỐNG]', style='Normal')
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER


    document.add_paragraph('• Backend (Server Side):', style='List Bullet')
    document.add_paragraph(
        '  Xây dựng bằng Node.js và Express.js. Chịu trách nhiệm xử lý logic nghiệp vụ, xác thực người dùng (JWT), '
        'và tương tác với cơ sở dữ liệu.', style='Normal'
    )
    document.add_paragraph('• Database:', style='List Bullet')
    document.add_paragraph(
        '  Sử dụng MongoDB - cơ sở dữ liệu NoSQL hướng tài liệu, phù hợp với cấu trúc dữ liệu linh hoạt của sản phẩm thuốc '
        'và đơn hàng.', style='Normal'
    )

    # 3.2 Components
    document.add_heading('3.2. Mô tả các thành phần trong hệ thống', level=2)
    
    # 3.2.1 Component Table
    table2 = document.add_table(rows=1, cols=2)
    table2.style = 'Table Grid'
    hdr_cells2 = table2.rows[0].cells
    hdr_cells2[0].text = 'Phân hệ'
    hdr_cells2[1].text = 'Chức năng chính'
    hdr_cells2[0].paragraphs[0].runs[0].font.bold = True
    hdr_cells2[1].paragraphs[0].runs[0].font.bold = True

    components = [
        ['Khách hàng (Customer)', 'Xem/Tìm kiếm thuốc, Giỏ hàng, Đặt hàng Online, Thanh toán VNPay, Xem lịch sử đơn'],
        ['Quản trị (Admin)', 'Quản lý toàn bộ hệ thống: User, Sản phẩm, Danh mục, Kho, Đơn hàng, Thống kê doanh thu'],
        ['Kho (Warehouse Staff)', 'Quản lý nhập kho (Import), Chuyển kho (Transfer), Kiểm kho, Quản lý lô hạn sử dụng'],
        ['Bán hàng (Sale Staff)', 'Bán hàng tại quầy (POS), Tạo đơn cho khách vãng lai, Tra cứu tồn kho']
    ]
    for comp in components:
        row = table2.add_row().cells
        row[0].text = comp[0]
        row[1].text = comp[1]

    # 3.3 Data Design
    document.add_heading('3.3. Thiết kế dữ liệu chi tiết', level=2)
    document.add_paragraph('Cơ sở dữ liệu MongoDB được thiết kế gồm các Collection chính, chia theo nhóm chức năng:')

    def add_collection_table(name, description, fields):
        document.add_heading(f'Collection: {name}', level=3)
        document.add_paragraph(f'Mô tả: {description}')
        
        table = document.add_table(rows=1, cols=3)
        table.style = 'Table Grid'
        hdr = table.rows[0].cells
        hdr[0].text = 'Tên trường (Field)'
        hdr[1].text = 'Kiểu dữ liệu (Type)'
        hdr[2].text = 'Ghi chú / Liên kết'
        for cell in hdr:
            cell.paragraphs[0].runs[0].font.bold = True
            
        for field in fields:
            row = table.add_row().cells
            row[0].text = field[0]
            row[1].text = field[1]
            row[2].text = field[2]
        document.add_paragraph() # Spacing

    # --- Group 1: Auth ---
    document.add_heading('3.3.1. Nhóm Quản lý Người dùng (Auth)', level=3)
    
    add_collection_table('User', 'Lưu trữ thông tin tài khoản (Admin, Staff, Customer).', [
        ['userName', 'String', 'Tên đăng nhập (Bắt buộc, Unique)'],
        ['passWord', 'String', 'Mật khẩu (hashed bcrypt)'],
        ['roleId', 'ObjectId', 'Liên kết Collection Role'],
        ['warehouseId', 'ObjectId', 'Liên kết Warehouse (Nơi làm việc)'],
        ['email', 'String', 'Email người dùng'],
        ['phoneNum', 'String', 'Số điện thoại'],
        ['isActive', 'Boolean', 'Trạng thái hoạt động (Default: true)']
    ])

    add_collection_table('Role', 'Định nghĩa các vai trò trong hệ thống.', [
        ['roleName', 'String', 'Tên vai trò (Admin, Staff, Customer, WarehouseStaff)'],
        ['description', 'String', 'Mô tả chi tiết quyền hạn']
    ])

    add_collection_table('OTP', 'Lưu mã xác thực tạm thời (Quên mật khẩu).', [
        ['email', 'String', 'Email nhận mã'],
        ['otp', 'String', 'Mã xác thực'],
        ['createAt', 'Date', 'Thời gian tạo (Tự xóa sau 5 phút)']
    ])

    # --- Group 2: Product ---
    document.add_heading('3.3.2. Nhóm Sản phẩm & Kho thuốc', level=3)

    add_collection_table('Product', 'Lưu trữ thông tin thuốc/sản phẩm.', [
        ['productName', 'String', 'Tên sản phẩm'],
        ['price', 'Number', 'Giá bán (Đơn vị cơ sở)'],
        ['stockQuantity', 'Number', 'Tổng tồn kho (Tính toán từ Batch)'],
        ['variants', 'Array', 'Quy đổi đơn vị (Vd: [Unit: "Vỉ", Price: 5000, Ratio: 10])'],
        ['categoryId', 'ObjectId', 'Liên kết Collection Category'],
        ['manufacturerId', 'ObjectId', 'Liên kết Collection Manufacturer'],
        ['nearestExpiryDate', 'Date', 'Ngày hết hạn gần nhất (để cảnh báo)']
    ])

    add_collection_table('ProductBatch', 'Quản lý lô sản xuất và hạn sử dụng (Thực thể kho chi tiết).', [
        ['batchNumber', 'String', 'Số lô sản xuất'],
        ['expiryDate', 'Date', 'Hạn sử dụng'],
        ['quantity', 'Number', 'Số lượng nhập ban đầu'],
        ['remainingQuantity', 'Number', 'Số lượng còn lại trong kho'],
        ['warehouseId', 'ObjectId', 'Kho chứa lô hàng này'],
        ['productId', 'ObjectId', 'Liên kết Product']
    ])

    add_collection_table('Category', 'Danh mục sản phẩm.', [
        ['categoryName', 'String', 'Tên danh mục thuốc']
    ])

    add_collection_table('Manufacturer', 'Nhà sản xuất/Nhà cung cấp.', [
        ['manufacturerName', 'String', 'Tên nhà sản xuất'],
        ['phoneNum', 'String', 'Số điện thoại liên hệ']
    ])

    # --- Group 3: Warehouse ---
    document.add_heading('3.3.3. Nhóm Quản lý Kho hàng', level=3)

    add_collection_table('Warehouse', 'Danh sách các kho hàng (Tổng/Chi nhánh).', [
        ['warehouseName', 'String', 'Tên kho'],
        ['warehouseType', 'String', 'Loại kho ("central" hoặc "branch")'],
        ['address', 'String', 'Địa chỉ kho'],
        ['manager', 'String', 'Tên quản lý kho']
    ])

    add_collection_table('InventoryTransfer', 'Phiếu chuyển kho (Giữa các chi nhánh).', [
        ['fromWarehouseId', 'ObjectId', 'Kho nguồn'],
        ['toWarehouseId', 'ObjectId', 'Kho đích'],
        ['products', 'Array', 'Danh sách SP chuyển: [{productId, quantity}]'],
        ['status', 'String', 'Trạng thái (Pending, Completed, Cancelled)']
    ])
    
    add_collection_table('PurchaseInvoice', 'Hóa đơn nhập hàng từ Nhà cung cấp.', [
        ['warehouseId', 'ObjectId', 'Kho nhập hàng'],
        ['manufacturerId', 'ObjectId', 'Nhà cung cấp'],
        ['dateImport', 'Date', 'Ngày nhập kho'],
        ['totalBill', 'Number', 'Tổng tiền nhập']
    ])

    # --- Group 4: Sales ---
    document.add_heading('3.3.4. Nhóm Bán hàng & Đơn hàng', level=3)

    add_collection_table('SaleInvoice', 'Hóa đơn bán hàng (Đơn hàng).', [
        ['userId', 'ObjectId', 'Khách hàng'],
        ['staffId', 'ObjectId', 'Nhân viên bán (Nếu bán tại quầy)'],
        ['totalAmount', 'Number', 'Tổng tiền thanh toán'],
        ['paymentMethod', 'String', 'Phương thức (VNPay, COD, Cash)'],
        ['statusId', 'ObjectId', 'Trạng thái đơn hàng (Ref: OrderStatus)'],
        ['isInStoreSale', 'Boolean', 'True: Bán tại quầy, False: Online'],
        ['vnpayTransactionNo', 'String', 'Mã giao dịch VNPay (Nếu có)']
    ])

    add_collection_table('SaleInvoiceDetail', 'Chi tiết sản phẩm trong đơn hàng.', [
        ['saleInvoiceId', 'ObjectId', 'Liên kết SaleInvoice'],
        ['productId', 'ObjectId', 'Sản phẩm'],
        ['batchId', 'ObjectId', 'Lấy từ lô hàng nào (Quản lý tồn chính xác)'],
        ['quantity', 'Number', 'Số lượng bán'],
        ['unitPrice', 'Number', 'Đơn giá bán']
    ])

    add_collection_table('OrderStatus', 'Trạng thái đơn hàng.', [
        ['statusName', 'String', 'Tên trạng thái (Pending, Shipping, Completed, Cancelled)']
    ])

    # --- Group 5: System ---
    document.add_heading('3.3.5. Nhóm Cấu hình Hệ thống', level=3)

    add_collection_table('Setting', 'Cấu hình chung cho Website.', [
        ['storeName', 'String', 'Tên cửa hàng hiển thị'],
        ['phone', 'String', 'Hotline liên hệ'],
        ['email', 'String', 'Email liên hệ'],
        ['notifyLowStock', 'Boolean', 'Bật/tắt cảnh báo sắp hết hàng']
    ])
    
    # Save
    file_path = 'Bao_Cao_Do_An.docx'
    document.save(file_path)
    print(f"Successfully created {file_path}")

if __name__ == '__main__':
    create_report()
